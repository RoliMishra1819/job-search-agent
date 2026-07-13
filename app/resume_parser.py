import os
import json
import pdfplumber
from docx import Document
from openai import OpenAI
from dotenv import load_dotenv

load_dotenv()
client = OpenAI(
    base_url="https://api.groq.com/openai/v1",
    api_key=os.getenv("GROQ_API_KEY"),
    max_retries=0,
)

RESUME_SCHEMA = {
    "name": "",
    "email": "",
    "phone": "",
    "location": "",
    "linkedin": "",
    "github": "",
    "summary": "",
    "skills": [],
    "experience": [],   # list of {title, company, duration, points: []}
    "projects": [],     # list of {name, description, tech: []}
    "education": [],    # list of {degree, institution, year}
    "certifications": [],
    "achievements": []
}


def extract_text_from_docx(path: str) -> str:
    """Read a .docx file and return its plain text content."""
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also pull text out of tables (some resumes use table layouts)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    lines.append(cell.text.strip())

    return "\n".join(lines)

def extract_text_from_pdf(path: str) -> str:
    """Read a PDF file and return its plain text content."""
    lines = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                lines.append(text)
    return "\n".join(lines)


def extract_text_from_file(path: str) -> str:
    """Dispatch to the right extractor based on file extension."""
    if path.lower().endswith(".pdf"):
        return extract_text_from_pdf(path)
    elif path.lower().endswith(".docx"):
        return extract_text_from_docx(path)
    else:
        raise ValueError(f"Unsupported file type: {path}")


def structure_resume(raw_text: str) -> dict:
    """Send raw resume text to the LLM and get back structured JSON."""
    prompt = f"""
You are a resume parser. Convert the following resume text into JSON
matching exactly this schema (keep keys, fill values, use [] or "" if unknown):

{json.dumps(RESUME_SCHEMA, indent=2)}

IMPORTANT RULES:
- The "skills" list must include EVERY skill line from the resume's Key Skills
  section, one array entry per category line. Do not skip, merge, or summarize any.
- Do not omit any project, education entry, or achievement present in the text.
- Copy skill/category lines verbatim rather than paraphrasing them.

Resume text:
\"\"\"
{raw_text}
\"\"\"

Return ONLY valid JSON. No markdown, no commentary.
"""

    for attempt in range(3):
        response = client.chat.completions.create(
            model="llama-3.1-8b-instant",
            messages=[{"role": "user", "content": prompt}],
            temperature=0,
            response_format={"type": "json_object"},
        )

        content = response.choices[0].message.content.strip()
        content = content.replace("```json", "").replace("```", "").strip()

        try:
            return json.loads(content)
        except json.JSONDecodeError:
            if attempt == 2:
                raise ValueError(f"LLM did not return valid JSON after 3 attempts:\n{content}")
            continue


def parse_resume(path: str) -> dict:
    """Full pipeline: docx/pdf -> raw text -> structured JSON."""
    raw_text = extract_text_from_file(path)
    if not raw_text.strip():
        raise ValueError(f"No text extracted from {path} — is the file empty?")
    return structure_resume(raw_text)


if __name__ == "__main__":
    result = parse_resume("resumes/master_resume.docx")
    print(json.dumps(result, indent=2))