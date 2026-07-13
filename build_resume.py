from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Name
name = doc.add_paragraph()
name_run = name.add_run("Roli Mishra")
name_run.bold = True
name_run.font.size = Pt(18)

# Contact line
contact = doc.add_paragraph("Kanpur, Uttar Pradesh | roli85756@gmail.com | 8869944959 | linkedin.com/in/roli-mishra")

# Summary
doc.add_heading("Summary", level=2)
doc.add_paragraph(
    "Data Analyst and AIML undergraduate with hands-on experience in SQL, Python, Excel, Tableau, "
    "and data visualization. Skilled in data cleaning, exploratory data analysis, dashboard creation, "
    "and predictive analytics. Passionate about transforming raw data into actionable business insights "
    "and supporting data-driven decision-making."
)

# Education
doc.add_heading("Education", level=2)
doc.add_paragraph("Dr APJ Abdul Kalam Technical University — B.Tech CS(AIML), 7.7 CGPA (09/2022 - 06/2026)")
doc.add_paragraph("Bright Angels Intermediate College — Intermediate, 71% (04/2020 - 03/2022)")
doc.add_paragraph("Bright Angels Intermediate College — High School, 80% (04/2018 - 03/2020)")

# Projects
doc.add_heading("Projects", level=2)

doc.add_paragraph("SafeSight AI – Smart Anomaly Detection System (01/2026 - 04/2026)", style="Heading 3")
for point in [
    "Built anomaly detection pipeline using YOLOv8n, ByteTrack & SimCLR on UCF Crime dataset",
    "Designed LSTM autoencoder + XGBoost for 5-level threat prediction (136-d features)",
    "Engineered behavioral features & risk scoring system for pre-incident alerts",
    "Deployed Streamlit dashboard with live analytics and SHAP explainability",
]:
    doc.add_paragraph(point, style="List Bullet")

doc.add_paragraph("Customer Churn Prediction (02/2025 - 03/2025)", style="Heading 3")
for point in [
    "Built a machine learning model to predict customer churn with 77% accuracy on 5K+ records",
    "Performed data preprocessing, feature engineering, and model evaluation",
    "Used Python libraries for data analysis and prediction",
]:
    doc.add_paragraph(point, style="List Bullet")

# Achievements
doc.add_heading("Achievements", level=2)
for point in [
    "Built a final-year project integrating Computer Vision, Deep Learning, and Explainable AI.",
    "Solved 200+ Python and SQL problems.",
    "Strong interest in applying machine learning and analytics to customer behavior and business optimization problems.",
]:
    doc.add_paragraph(point, style="List Bullet")

# Key Skills
doc.add_heading("Key Skills", level=2)
skills = [
    "SQL: Joins, Subqueries, Window Functions, Aggregations, Data Extraction",
    "Excel: Pivot Tables, Lookup Functions, Data Cleaning, Reporting, Dashboards",
    "Data Visualization: Tableau, Matplotlib, Dashboard Development, Data Storytelling",
    "Python: Pandas, NumPy, Data Analysis, Data Manipulation",
    "Analytics: EDA, Statistical Analysis, KPI Analysis, Customer Analytics, Business Metrics, Trend Analysis",
    "Databases: MySQL",
    "Machine Learning: Regression, Classification, XGBoost, Decision Trees, Predictive Analytics",
    "Tools: Jupyter Notebook, Git, GitHub, Streamlit",
    "Concepts: DBMS, OOP, SDLC",
]
for s in skills:
    doc.add_paragraph(s, style="List Bullet")

# Interests
doc.add_heading("Interests", level=2)
for i in ["Coding", "Playing chess", "Learning new programming languages"]:
    doc.add_paragraph(i, style="List Bullet")

doc.save("resumes/master_resume.docx")
print("Saved resumes/master_resume.docx")