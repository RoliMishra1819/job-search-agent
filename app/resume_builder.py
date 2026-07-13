import os
import re
from docx import Document
from docx.shared import Pt

from resume_parser import parse_resume
from jd_parser import parse_jd
from matcher import match_resume_to_jd


def sanitize_filename(name: str) -> str:
    """Make a company name safe for use in a filename."""
    name = re.sub(r"[^\w\s-]", "", name).strip()
    return re.sub(r"[\s]+", "_", name)


def build_tailored_resume(resume_json: dict, match_result: dict, company_name: str, output_dir: str = "resumes/generated") -> str:
    """
    Generate a tailored resume docx: reorders skills so matched/required
    skills appear first, keeps all original content (truthful, nothing invented).
    """
    os.makedirs(output_dir, exist_ok=True)

    matched = set(s.lower() for s in match_result.get("matched_required", []) + match_result.get("matched_nice_to_have", []))

    original_skills = resume_json.get("skills", [])

    def is_matched(skill_line: str) -> bool:
        return any(m in skill_line.lower() for m in matched)

    # Matched-relevant skill lines first, rest after — nothing removed, nothing invented
    prioritized_skills = sorted(original_skills, key=lambda s: not is_matched(s))

    doc = Document()

    name = doc.add_paragraph()
    name_run = name.add_run(resume_json.get("name", ""))
    name_run.bold = True
    name_run.font.size = Pt(18)

    contact_parts = [
        resume_json.get("location", ""),
        resume_json.get("email", ""),
        resume_json.get("phone", ""),
        resume_json.get("linkedin", ""),
        resume_json.get("github", ""),
    ]
    doc.add_paragraph(" | ".join(p for p in contact_parts if p))

    doc.add_heading("Summary", level=2)
    doc.add_paragraph(resume_json.get("summary", ""))

    doc.add_heading("Key Skills", level=2)
    for s in prioritized_skills:
        doc.add_paragraph(s, style="List Bullet")

    doc.add_heading("Projects", level=2)
    for proj in resume_json.get("projects", []):
        doc.add_paragraph(proj.get("name", ""), style="Heading 3")
        desc = proj.get("description", "")
        if desc:
            doc.add_paragraph(desc, style="List Bullet")

    doc.add_heading("Experience", level=2)
    for exp in resume_json.get("experience", []):
        title_line = f"{exp.get('role', '')} ({exp.get('duration', '')})"
        doc.add_paragraph(title_line, style="Heading 3")
        desc = exp.get("description", "")
        if desc:
            doc.add_paragraph(desc, style="List Bullet")

    doc.add_heading("Education", level=2)
    for edu in resume_json.get("education", []):
        line = f"{edu.get('degree', '')} — {edu.get('institution', '')} ({edu.get('duration', '')})"
        doc.add_paragraph(line)

    filename = f"Resume_{sanitize_filename(company_name)}.docx"
    output_path = os.path.join(output_dir, filename)
    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    resume_data = parse_resume("resumes/master_resume.docx")

    sample_jd = """
    Data Analyst - Fresher
    Company: Example Corp
    Location: Bangalore, India

    We are looking for a fresher Data Analyst with strong skills in:
    - Python
    - SQL
    - Power BI
    - Excel
    - Tableau
    - Statistics

    Nice to have: AWS, Azure
    """
    jd_data = parse_jd(sample_jd)
    match_result = match_resume_to_jd(resume_data, jd_data)

    path = build_tailored_resume(resume_data, match_result, jd_data.get("company", "UnknownCompany"))
    print(f"Tailored resume saved to: {path}")